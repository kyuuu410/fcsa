"""Offline tests for the public Google Drive record synchronizer."""

from collections import deque
from datetime import datetime
import hashlib
import io
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import Mock, patch
from urllib.error import HTTPError, URLError

import openpyxl

import sync_records as sync


ENVIRONMENT = {
    "FCSA_DRIVE_FILE_ID": "valid_drive_file_id_123",
}
LAST_MODIFIED = "Mon, 07 Sep 2026 06:34:39 GMT"


def workbook_fixture(*, player_name="테스트 선수", wins=1, cached_formula=True):
    workbook = openpyxl.Workbook()
    roster = workbook.active
    roster.title = "26년 선수 기록"
    team = workbook.create_sheet("26년 팀 기록")
    attendance = workbook.create_sheet("26년 출장 기록")
    goals = workbook.create_sheet("26년 득점 기록")
    assists = workbook.create_sheet("26년 도움 기록")

    roster.append(["포지션", "등번호", "이름", "득점", "도움", "공격포인트", "출장횟수", "참석율"])
    roster.append(["FW", 9, player_name, 1 if cached_formula else "=1", 0, 1, 1, 1])

    team["A2"], team["B2"] = "전적", 1
    team["A3"], team["B3"] = "승리", wins
    team["A4"], team["B4"] = "무승부", 0
    team["A5"], team["B5"] = "패배", 1 - wins
    team["A6"], team["B6"] = "승률", wins
    team["A7"], team["B7"] = "득점", 1
    team["A8"], team["B8"] = "실점", 0
    for address, value in {"D1": "vs", "E1": "승", "F1": "무", "G1": "패", "H1": "승률"}.items():
        team[address] = value
    team["D2"], team["E2"], team["F2"], team["G2"], team["H2"] = "상대 FC", wins, 0, 1 - wins, wins

    match_date = datetime(2026, 1, 3)
    for sheet, event_value in ((attendance, "O"), (goals, 1), (assists, None)):
        sheet["A1"], sheet["B1"], sheet["C1"] = "포지션", "이름", match_date
        sheet["A2"], sheet["B2"], sheet["C2"] = "FW", player_name, event_value
        sheet["A3"], sheet["C3"] = "합계", 1 if sheet is not assists else 0

    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def write_records(path, sha256, *, title="승인된 제목.xlsx", imported_at="unchanged-time"):
    records = {
        "source": {"title": title, "sha256": sha256, "importedAt": imported_at},
        "players": [{"name": "기존 선수"}],
        "dates": [{}],
        "qualityNotes": [],
    }
    path.write_text(f"window.FCSA_RECORDS = {json.dumps(records, ensure_ascii=False)};\n", encoding="utf-8")
    return records


class FakeResponse:
    def __init__(self, url, content, headers=None, status=200):
        self.url = url
        self.content = content
        self.status = status
        self.headers = {
            "Content-Length": str(len(content)),
            "Content-Type": "application/octet-stream",
            **(headers or {}),
        }

    def __enter__(self):
        return self

    def __exit__(self, *_args):
        return False

    def geturl(self):
        return self.url

    def read(self, limit):
        return self.content[:limit]


class FakeOpener:
    def __init__(self, responses):
        self.responses = deque(responses)
        self.requests = []

    def open(self, request, timeout):
        self.requests.append((request, timeout))
        response = self.responses.popleft()
        if isinstance(response, BaseException):
            raise response
        return response


class RecordSyncTests(unittest.TestCase):
    def response(self, content, *, url="https://drive.usercontent.google.com/download", headers=None):
        return FakeResponse(url, content, {"Last-Modified": LAST_MODIFIED, **(headers or {})})

    def run_sync(self, output, content, *, headers=None):
        opener = FakeOpener([self.response(content, headers=headers)])
        return sync.sync_records(output, environment=ENVIRONMENT, opener=opener, sleep=lambda _: None)

    def test_changed_workbook_updates_member_and_team_records(self):
        content = workbook_fixture(player_name="새 선수", wins=0)
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "records.js"
            write_records(output, "0" * 64)
            result = self.run_sync(output, content)
            records = sync.read_existing_records(output)

        self.assertTrue(result.changed)
        self.assertEqual(records["players"][0]["name"], "새 선수")
        self.assertEqual(records["summary"]["losses"], 1)
        self.assertEqual(records["opponents"][0]["losses"], 1)
        self.assertEqual(records["source"]["title"], "승인된 제목.xlsx")
        self.assertEqual(records["source"]["sha256"], hashlib.sha256(content).hexdigest())

    def test_public_transport_updates_then_downloads_and_short_circuits(self):
        content = workbook_fixture(player_name="전송 경로 선수")
        opener = FakeOpener([self.response(content), self.response(content)])
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "records.js"
            changed = sync.sync_records(output, environment=ENVIRONMENT, opener=opener, sleep=lambda _: None)
            original = output.read_bytes()
            with patch.object(sync, "import_records") as import_mock:
                unchanged = sync.sync_records(output, environment=ENVIRONMENT, opener=opener, sleep=lambda _: None)
            records = sync.read_existing_records(output)
            self.assertEqual(output.read_bytes(), original)

        self.assertTrue(changed.changed)
        self.assertFalse(unchanged.changed)
        self.assertEqual(records["players"][0]["name"], "전송 경로 선수")
        self.assertEqual(records["source"]["modifiedAt"], "2026-09-07T06:34:39+00:00")
        import_mock.assert_not_called()
        self.assertEqual(len(opener.requests), 2)
        for request, timeout in opener.requests:
            self.assertEqual(request.get_method(), "GET")
            self.assertEqual(request.full_url, f"https://drive.google.com/uc?export=download&id={ENVIRONMENT['FCSA_DRIVE_FILE_ID']}")
            self.assertIsNone(request.get_header("Authorization"))
            self.assertIsNone(request.get_header("Cookie"))
            self.assertEqual(request.get_header("Cache-control"), "no-cache")
            self.assertEqual(request.get_header("Pragma"), "no-cache")
            self.assertEqual(timeout, sync.TIMEOUT_SECONDS)

    def test_missing_or_invalid_last_modified_is_preserved_as_unknown(self):
        content = workbook_fixture()
        cases = [None, "not a date"]
        for value in cases:
            headers = {"Last-Modified": value} if value is not None else {}
            with self.subTest(value=value), tempfile.TemporaryDirectory() as directory:
                output = Path(directory) / "records.js"
                opener = FakeOpener([FakeResponse("https://drive.usercontent.google.com/download", content, headers)])
                sync.sync_records(output, environment=ENVIRONMENT, opener=opener, sleep=lambda _: None)
                records = sync.read_existing_records(output)
                self.assertIsNone(records["source"]["modifiedAt"])

    def test_rejects_unapproved_download_locations(self):
        allowed = [
            "https://drive.google.com/uc?export=download&id=x",
            "https://drive.usercontent.google.com/download?id=x",
        ]
        rejected = [
            "http://drive.google.com/uc?id=x",
            "https://drive.google.com.evil.example/uc?id=x",
            "https://user@drive.google.com/uc?id=x",
            "https://drive.google.com:444/uc?id=x",
        ]
        for url in allowed:
            with self.subTest(url=url):
                self.assertEqual(sync._validate_download_url(url), url)
        for url in rejected:
            with self.subTest(url=url), self.assertRaisesRegex(sync.RecordSyncError, "unapproved"):
                sync._validate_download_url(url)

        opener = FakeOpener([self.response(workbook_fixture(), url="https://example.invalid/download")])
        with tempfile.TemporaryDirectory() as directory, self.assertRaisesRegex(sync.RecordSyncError, "unapproved"):
            sync.sync_records(Path(directory) / "records.js", environment=ENVIRONMENT, opener=opener, sleep=lambda _: None)

    def test_rejects_html_partial_oversized_and_incomplete_responses(self):
        content = workbook_fixture()
        cases = [
            (FakeResponse("https://drive.google.com/uc", b"<html>login</html>", {"Content-Type": "text/html"}), "XLSX download"),
            (FakeResponse("https://drive.usercontent.google.com/download", content, status=206), "complete download"),
            (FakeResponse("https://drive.usercontent.google.com/download", content, {"Content-Length": str(sync.MAX_XLSX_BYTES + 1)}), "size limit"),
            (FakeResponse("https://drive.usercontent.google.com/download", content, {"Content-Length": str(len(content) + 1)}), "incomplete response"),
        ]
        for response, message in cases:
            with self.subTest(message=message), tempfile.TemporaryDirectory() as directory:
                output = Path(directory) / "records.js"
                write_records(output, "1" * 64)
                original = output.read_bytes()
                with self.assertRaisesRegex(sync.RecordSyncError, message):
                    sync.sync_records(output, environment=ENVIRONMENT, opener=FakeOpener([response]), sleep=lambda _: None)
                self.assertEqual(output.read_bytes(), original)

    def test_bad_archive_preserves_existing_output(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "records.js"
            write_records(output, "1" * 64)
            original = output.read_bytes()
            with self.assertRaisesRegex(sync.RecordSyncError, "archive structure"):
                self.run_sync(output, b"not an xlsx archive")
            self.assertEqual(output.read_bytes(), original)

    def test_formula_without_cached_value_is_rejected_and_preserves_output(self):
        content = workbook_fixture(cached_formula=False)
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "records.js"
            write_records(output, "2" * 64)
            original = output.read_bytes()
            with self.assertRaisesRegex(sync.RecordSyncError, "validation or import failed"):
                self.run_sync(output, content)
            self.assertEqual(output.read_bytes(), original)

    def test_network_failures_retry_without_exposing_details(self):
        file_id = ENVIRONMENT["FCSA_DRIVE_FILE_ID"]
        secret_detail = "secret@example.com/path"
        cases = [
            (HTTPError("https://private.example/secret", 503, secret_detail, {}, None), "failed with HTTP 503", sync.MAX_ATTEMPTS),
            (URLError(secret_detail), "failed after retries", sync.MAX_ATTEMPTS),
            (HTTPError("https://private.example/secret", 403, secret_detail, {}, None), "HTTP 403", 1),
        ]
        for failure, expected, attempts in cases:
            opener = Mock()
            opener.open.side_effect = failure
            sleeps = []
            try:
                with self.subTest(expected=expected), self.assertRaises(sync.RecordSyncError) as caught:
                    sync.download_xlsx(file_id, opener=opener, sleep=sleeps.append)
                message = str(caught.exception)
                self.assertIn(expected, message)
                self.assertNotIn(file_id, message)
                self.assertNotIn(secret_detail, message)
                self.assertEqual(opener.open.call_count, attempts)
                request = opener.open.call_args.args[0]
                self.assertEqual(request.get_method(), "GET")
                self.assertIsNone(request.get_header("Authorization"))
                self.assertEqual(opener.open.call_args.kwargs["timeout"], sync.TIMEOUT_SECONDS)
            finally:
                if isinstance(failure, HTTPError):
                    failure.close()

    def test_rejects_invalid_file_ids_and_writes_github_output(self):
        for value in (None, "short", "contains/slash", "contains whitespace"):
            with self.subTest(value=value), self.assertRaisesRegex(sync.RecordSyncError, "missing or invalid"):
                sync.sync_records(environment={"FCSA_DRIVE_FILE_ID": value}, opener=Mock())
        with tempfile.TemporaryDirectory() as directory:
            destination = Path(directory) / "github-output"
            sync.write_github_output(True, {"GITHUB_OUTPUT": str(destination)})
            sync.write_github_output(False, {"GITHUB_OUTPUT": str(destination)})
            self.assertEqual(destination.read_text(encoding="utf-8"), "changed=true\nchanged=false\n")


if __name__ == "__main__":
    unittest.main()
