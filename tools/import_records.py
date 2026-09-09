"""Read the FC SSA workbook without modifying it and export validated website data."""

import argparse
import hashlib
import io
import json
import math
import os
from pathlib import Path
import tempfile
from datetime import date, datetime, timezone

import openpyxl


class RecordValidationError(ValueError):
    pass


ROLES = {"GK": "골키퍼", "DF": "수비수", "MF": "미드필더", "FW": "공격수"}
METRIC_LABELS = {"goals": "득점", "assists": "도움", "appearances": "출장횟수", "participants": "참여 인원", "wins": "승리", "draws": "무승부", "losses": "패배"}


def location(sheet, address):
    return f"'{sheet.title}'!{address}"


def require_value(sheet, address):
    value = sheet[address].value
    if value is None or (isinstance(value, str) and not value.strip()):
        raise RecordValidationError(f"Missing value: {location(sheet, address)}")
    if sheet[address].data_type == "e":
        raise RecordValidationError(f"Formula error: {location(sheet, address)}")
    return value


def number(sheet, address, integer=True):
    value = require_value(sheet, address)
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise RecordValidationError(f"Expected number: {location(sheet, address)}")
    if not math.isfinite(value) or value < 0 or (integer and value != int(value)):
        raise RecordValidationError(f"Invalid number: {location(sheet, address)}")
    return int(value) if integer else float(value)


def header(sheet, address, expected):
    if require_value(sheet, address) != expected:
        raise RecordValidationError(f"Unexpected header: {location(sheet, address)}; expected {expected!r}")


def event_count(sheet, address):
    # In these sparse event matrices an empty cell means no recorded event.
    return 0 if sheet[address].value is None else number(sheet, address)


def convert_workbook(content, source_url, modified_at, source_title):
    workbook = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    try:
        required_sheets = [f"26년 {kind} 기록" for kind in ["선수", "팀", "출장", "득점", "도움"]]
        if any(name not in workbook.sheetnames for name in required_sheets):
            raise RecordValidationError("Missing required 2026 record sheet")
        roster, team, attendance, goals, assists = [workbook[name] for name in required_sheets]
        for column, name in enumerate(["포지션", "등번호", "이름", "득점", "도움", "공격포인트", "출장횟수", "참석율"], 1):
            header(roster, f"{openpyxl.utils.get_column_letter(column)}1", name)
        for sheet in [attendance, goals, assists]:
            header(sheet, "A1", "포지션")
            header(sheet, "B1", "이름")
        for address, name in {"A2": "전적", "A3": "승리", "A4": "무승부", "A5": "패배", "A6": "승률", "A7": "득점", "A8": "실점", "D1": "vs", "E1": "승", "F1": "무", "G1": "패", "H1": "승률"}.items():
            header(team, address, name)
        summary = {key: number(team, f"B{row}", key != "winRate") for key, row in {"played": 2, "wins": 3, "draws": 4, "losses": 5, "winRate": 6, "goalsFor": 7, "goalsAgainst": 8}.items()}
        if summary["played"] <= 0 or summary["winRate"] > 1:
            raise RecordValidationError("Invalid season summary")
        notes = []
        row = 2
        while roster.cell(row, 3).value is not None:
            row += 1
        end_row = row - 1
        if end_row < 2 or any(roster.cell(r, 3).value is not None for r in range(row, roster.max_row + 1)):
            raise RecordValidationError("Missing or discontinuous player rows")
        total_row = end_row + 1
        for sheet in [attendance, goals, assists]:
            header(sheet, f"A{total_row}", "합계")
        dates = []
        date_columns = []
        seen_dates = set()
        for column in range(3, max(sheet.max_column for sheet in [attendance, goals, assists]) + 1):
            address = f"{openpyxl.utils.get_column_letter(column)}1"
            value = require_value(attendance, address)
            if not isinstance(value, (date, datetime)) or value.year != 2026:
                raise RecordValidationError(f"Invalid date: {location(attendance, address)}")
            iso_date = value.strftime("%Y-%m-%d")
            if iso_date in seen_dates or (dates and iso_date <= dates[-1]["date"]):
                raise RecordValidationError("Dates must be unique and increasing")
            seen_dates.add(iso_date)
            for sheet in [goals, assists]:
                if require_value(sheet, address) != value:
                    raise RecordValidationError(f"Date mismatch: {location(sheet, address)}")
            if value.weekday() != 5:
                notes.append(f"{location(attendance, address)}의 {iso_date}는 토요일이 아닙니다. 원본 날짜를 유지했습니다.")
            total_address = f"{openpyxl.utils.get_column_letter(column)}{total_row}"
            dates.append({"date": iso_date, "participants": number(attendance, total_address), "goals": number(goals, total_address), "assists": number(assists, total_address)})
            date_columns.append(column)
        if not dates:
            raise RecordValidationError("No dated records")
        players = []
        seen_numbers = set()
        seen_names = set()
        for row in range(2, end_row + 1):
            position = require_value(roster, f"A{row}")
            name = require_value(roster, f"C{row}")
            shirt_number = number(roster, f"B{row}")
            if position not in ROLES or not isinstance(name, str) or name in seen_names or shirt_number in seen_numbers:
                raise RecordValidationError(f"Invalid or duplicate player: row {row}")
            seen_numbers.add(shirt_number)
            seen_names.add(name)
            for sheet in [attendance, goals, assists]:
                header(sheet, f"A{row}", position)
                header(sheet, f"B{row}", name)
            player = {"id": f"player-{shirt_number}", "name": name, "number": shirt_number, "position": position, "role": ROLES[position]}
            player.update({key: number(roster, f"{column}{row}", key != "attendanceRate") for key, column in {"goals": "D", "assists": "E", "points": "F", "appearances": "G", "attendanceRate": "H"}.items()})
            if player["attendanceRate"] > 1:
                raise RecordValidationError(f"Invalid attendance rate: row {row}")
            records = []
            for index, column in enumerate(date_columns):
                address = f"{openpyxl.utils.get_column_letter(column)}{row}"
                attendance_value = attendance[address].value
                attended = True if attendance_value == "O" else False if attendance_value is None else None
                if attended is None:
                    notes.append(f"{location(attendance, address)}에 O 또는 빈 셀이 아닌 값이 있습니다. 출장 여부는 확인 불가로 보존했습니다.")
                record = {"date": dates[index]["date"], "attended": attended, "goals": event_count(goals, address), "assists": event_count(assists, address)}
                if (record["goals"] or record["assists"]) and attended is not True:
                    notes.append(f"{name}의 {record['date']} 득점·도움 기록과 출장 표시가 일치하지 않습니다 ({location(attendance, address)}, {location(goals, address)}, {location(assists, address)}).")
                records.append(record)
            player["matchRecords"] = records
            for key in ["goals", "assists", "appearances"]:
                detail = sum(record["attended"] is True for record in records) if key == "appearances" else sum(record[key] for record in records)
                if detail != player[key]:
                    notes.append(f"{name}의 {METRIC_LABELS[key]} 공식 합계({player[key]})와 날짜별 합계({detail})가 다릅니다 ('26년 선수 기록' 행 {row}).")
            if player["points"] != player["goals"] + player["assists"]:
                notes.append(f"{location(roster, f'F{row}')} 공격포인트와 득점·도움 합계가 다릅니다.")
            if not math.isclose(player["attendanceRate"], player["appearances"] / summary["played"], abs_tol=1e-9):
                notes.append(f"{location(roster, f'H{row}')} 참석율과 출장횟수/전적이 다릅니다.")
            players.append(player)
        for index, record in enumerate(dates):
            column_letter = openpyxl.utils.get_column_letter(date_columns[index])
            for key, sheet in [("participants", attendance), ("goals", goals), ("assists", assists)]:
                detail = sum(player["matchRecords"][index]["attended"] is True for player in players) if key == "participants" else sum(player["matchRecords"][index][key] for player in players)
                if detail != record[key]:
                    notes.append(f"{record['date']} {METRIC_LABELS[key]} 원본 합계({record[key]})와 상세 합계({detail})가 다릅니다 ({location(sheet, f'{column_letter}{total_row}')}).")
            if record["assists"] > record["goals"]:
                notes.append(f"{record['date']} 기록은 득점 {record['goals']}, 도움 {record['assists']}입니다. 원본 값을 유지했습니다.")
        opponents = []
        opponent_names = set()
        for row in range(2, team.max_row + 1):
            if all(team.cell(row, column).value is None for column in range(4, 9)):
                continue
            name = require_value(team, f"D{row}")
            if not isinstance(name, str) or name in opponent_names:
                raise RecordValidationError(f"Invalid or duplicate opponent: row {row}")
            opponent_names.add(name)
            opponent = {"name": name, **{key: number(team, f"{column}{row}", key != "winRate") for key, column in {"wins": "E", "draws": "F", "losses": "G", "winRate": "H"}.items()}}
            played = opponent["wins"] + opponent["draws"] + opponent["losses"]
            if played <= 0 or opponent["winRate"] > 1:
                raise RecordValidationError(f"Invalid opponent record: row {row}")
            if not math.isclose(opponent["winRate"], opponent["wins"] / played, abs_tol=1e-9):
                notes.append(f"{location(team, f'H{row}')} 승률과 상대별 전적이 다릅니다.")
            opponents.append(opponent)
        for key in ["wins", "draws", "losses"]:
            if sum(opponent[key] for opponent in opponents) != summary[key]:
                notes.append(f"상대별 {METRIC_LABELS[key]} 합계와 원본 시즌 합계({summary[key]})가 다릅니다.")
        if sum(summary[key] for key in ["wins", "draws", "losses"]) != summary["played"]:
            notes.append("원본 시즌 승·무·패 합계와 전적이 다릅니다.")
        if len(dates) != summary["played"]:
            notes.append(f"원본 전적 {summary['played']}경기와 날짜 열 {len(dates)}개가 다릅니다.")
        if not math.isclose(summary["winRate"], summary["wins"] / summary["played"], abs_tol=1e-9):
            notes.append("원본 시즌 승률과 승리/전적이 다릅니다.")
        detailed_goals = sum(player["goals"] for player in players)
        if detailed_goals != summary["goalsFor"]:
            notes.append(f"공식 팀 득점은 {summary['goalsFor']}골('26년 팀 기록'!B7), 선수별 득점 합계는 {detailed_goals}골('26년 선수 기록'!D2:D{end_row})입니다. 차이를 보정하지 않았습니다.")
        notes.append("날짜별 상대팀·실점·경기 결과·쿼터 기록은 원본에 없어 제공하지 않습니다.")
        return {"version": 1, "season": 2026, "source": {"title": source_title, "modifiedAt": modified_at, "importedAt": datetime.now(timezone.utc).isoformat(), "sha256": hashlib.sha256(content).hexdigest()}, "summary": summary, "players": players, "dates": dates, "opponents": opponents, "qualityNotes": notes}
    finally:
        workbook.close()


def import_records(source, output, source_url, modified_at, source_title="2026_FC쏘아 스탯.xlsx"):
    source = Path(source).resolve()
    output = Path(output).resolve()
    if source == output:
        raise RecordValidationError("Output must not overwrite the source workbook")
    content = source.read_bytes()
    records = convert_workbook(content, source_url, modified_at, source_title)
    serialized = json.dumps(records, ensure_ascii=False, allow_nan=False, separators=(",", ":"))
    # Also safe if this external script is later embedded in an HTML script tag.
    for character, escaped in [("<", "\\u003c"), (">", "\\u003e"), ("&", "\\u0026"), ("\u2028", "\\u2028"), ("\u2029", "\\u2029")]:
        serialized = serialized.replace(character, escaped)
    if hashlib.sha256(source.read_bytes()).hexdigest() != records["source"]["sha256"]:
        raise RecordValidationError("Source changed while importing; output was not replaced")
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary_path = None
    try:
        with tempfile.NamedTemporaryFile(mode="w", encoding="utf-8", newline="\n", dir=output.parent, suffix=".tmp", delete=False) as temporary:
            temporary_path = Path(temporary.name)
            temporary.write(f"window.FCSA_RECORDS = {serialized};\n")
            temporary.flush()
            os.fsync(temporary.fileno())
        if hashlib.sha256(source.read_bytes()).hexdigest() != records["source"]["sha256"]:
            raise RecordValidationError("Source changed before publication; output was not replaced")
        os.replace(temporary_path, output)
        temporary_path = None
    finally:
        if temporary_path is not None:
            temporary_path.unlink(missing_ok=True)
    return records


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--modified-at", required=True)
    parser.add_argument("--source-title", default="2026_FC쏘아 스탯.xlsx")
    args = parser.parse_args()
    try:
        records = import_records(args.source, args.output, None, args.modified_at, args.source_title)
    except (OSError, ValueError, KeyError) as error:
        parser.exit(1, f"Import failed: {error}\n")
    print(f"Imported {len(records['players'])} players, {len(records['dates'])} dates, {len(records['qualityNotes'])} quality notes. Source SHA-256: {records['source']['sha256']}")


if __name__ == "__main__":
    main()
